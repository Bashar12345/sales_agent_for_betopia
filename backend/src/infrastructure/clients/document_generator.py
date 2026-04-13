# [OWNER: Asif — P2 Requirements Agent]
"""DocumentGenerator — implements IDocumentGeneratorPort using python-docx.

Generates a formatted .docx quotation file and saves it to DOC_OUTPUT_DIR.
The file path is stored on the Quotation record so the API can serve it.
"""

import os
import uuid
from datetime import datetime, timezone
from pathlib import Path

import structlog
from docx import Document
from docx.shared import Pt, RGBColor

from src.application.ports.document_generator_port import IDocumentGeneratorPort
from src.core.exceptions import DocumentGenerationError
from src.core.settings import settings
from src.domain.entities.quotation import Quotation

log = structlog.get_logger()


class DocumentGenerator(IDocumentGeneratorPort):
    def __init__(self) -> None:
        Path(settings.DOC_OUTPUT_DIR).mkdir(parents=True, exist_ok=True)

    async def generate_quotation_doc(
        self,
        quotation: Quotation,
        lead_name: str,
        agent_name: str,
    ) -> str:
        try:
            doc = Document()

            # ── Title ─────────────────────────────────────────────────────────
            title_para = doc.add_heading(quotation.title, level=1)
            title_para.runs[0].font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)

            # ── Meta ──────────────────────────────────────────────────────────
            doc.add_paragraph(f"Prepared for: {lead_name}")
            doc.add_paragraph(f"Prepared by: {agent_name}")
            doc.add_paragraph(
                f"Date: {datetime.now(timezone.utc).strftime('%B %d, %Y')}"
            )
            doc.add_paragraph()

            # ── Line items table ──────────────────────────────────────────────
            table = doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            header = table.rows[0].cells
            for idx, text in enumerate(["Description", "Qty", "Unit Price", "Total"]):
                header[idx].text = text
                header[idx].paragraphs[0].runs[0].bold = True

            for item in quotation.line_items:
                row = table.add_row().cells
                row[0].text = item.description
                row[1].text = str(item.quantity)
                row[2].text = f"{quotation.currency} {item.unit_price:,.2f}"
                row[3].text = f"{quotation.currency} {item.total:,.2f}"

            doc.add_paragraph()

            # ── Total ─────────────────────────────────────────────────────────
            total_para = doc.add_paragraph()
            total_run = total_para.add_run(
                f"Total: {quotation.currency} {quotation.total_amount:,.2f}"
            )
            total_run.bold = True
            total_run.font.size = Pt(14)

            # ── Notes ─────────────────────────────────────────────────────────
            if quotation.notes:
                doc.add_heading("Notes", level=2)
                doc.add_paragraph(quotation.notes)

            # ── Save ──────────────────────────────────────────────────────────
            filename = f"quotation_{quotation.id}_{uuid.uuid4().hex[:6]}.docx"
            output_path = os.path.join(settings.DOC_OUTPUT_DIR, filename)
            doc.save(output_path)
            log.info("doc.generated", path=output_path)
            return output_path
        except Exception as exc:
            raise DocumentGenerationError(f"Failed to generate .docx: {exc}") from exc
